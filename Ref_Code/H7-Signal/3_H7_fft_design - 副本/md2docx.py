#!/usr/bin/env python3
"""
将 Markdown 文件转换为 Word (.docx) 文档。
支持：标题(h1-h4)、代码块、表格、加粗、行内代码、列表、分割线、引用块
"""

import re
import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── 样式配置 ──────────────────────────────────────────
CODE_FONT_NAME = "Consolas"
CODE_FONT_SIZE = Pt(9)
BODY_FONT_NAME = "微软雅黑"
BODY_FONT_SIZE = Pt(11)
HEADING_FONT_NAME = "微软雅黑"

# ── 读取 Markdown ──────────────────────────────────────
SRC = r"D:\stm\学习\STM32H7_signal\3_H7_fft_design\DSP核心处理链深度分析.md"
DST = r"D:\stm\学习\STM32H7_signal\3_H7_fft_design\DSP核心处理链深度分析.docx"

with open(SRC, "r", encoding="utf-8") as f:
    lines = f.readlines()

doc = Document()

# ── 页面设置 ───────────────────────────────────────────
for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

# ── 默认段落样式 ──────────────────────────────────────
style = doc.styles['Normal']
font = style.font
font.name = BODY_FONT_NAME
font.size = BODY_FONT_SIZE
style.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT_NAME)

# ── 辅助函数 ───────────────────────────────────────────

def add_heading_styled(text, level):
    """添加标题，使用中文字体"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = HEADING_FONT_NAME
        run.element.rPr.rFonts.set(qn('w:eastAsia'), HEADING_FONT_NAME)
    return h

def add_code_block(code_lines, language=""):
    """添加带背景色的代码块"""
    code_text = "".join(code_lines)
    # 去掉末尾多余的换行
    code_text = code_text.rstrip('\n\r')

    for i, line in enumerate(code_text.split('\n')):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(14)
        p.paragraph_format.left_indent = Cm(0.3)

        # 灰色背景
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F0F0F0" w:val="clear"/>')
        p.paragraph_format.element.get_or_add_pPr().append(shading)

        if line == "":
            run = p.add_run(" ")
        else:
            run = p.add_run(line)
        run.font.name = CODE_FONT_NAME
        run.font.size = CODE_FONT_SIZE
        run.element.rPr.rFonts.set(qn('w:eastAsia'), CODE_FONT_NAME)

    # 代码块后加一个空行
    doc.add_paragraph()

def add_table_from_md(table_lines):
    """解析 Markdown 表格并添加到文档"""
    # 第一行是表头，第二行是分隔符 |---|，后面是数据行
    if len(table_lines) < 2:
        return

    # 解析表头
    header_cells = [c.strip() for c in table_lines[0].split('|')[1:-1]]
    if not header_cells:
        return

    # 跳过第二行（分隔符）
    data_rows = []
    for line in table_lines[2:]:
        cells = [c.strip() for c in line.split('|')[1:-1]]
        if cells:
            data_rows.append(cells)

    num_cols = len(header_cells)
    num_rows = 1 + len(data_rows)

    table = doc.add_table(rows=num_rows, cols=num_cols, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for j, text in enumerate(header_cells):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.name = BODY_FONT_NAME
        run.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT_NAME)
        # 蓝色背景
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="4472C4" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # 数据行
    for i, row in enumerate(data_rows):
        for j, text in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            run.font.name = BODY_FONT_NAME
            run.element.rPr.rFonts.set(qn('w:eastAsia'), BODY_FONT_NAME)
            # 交替行背景
            if i % 2 == 1:
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8F0FE" w:val="clear"/>')
                cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph()

def add_runs_to_paragraph(p, text):
    """解析行内格式（**加粗**、`代码`）到已有段落"""
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if token.startswith('**') and token.endswith('**'):
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('`') and token.endswith('`'):
            run = p.add_run(token[1:-1])
            run.font.name = CODE_FONT_NAME
            run.font.size = CODE_FONT_SIZE
            run.element.rPr.rFonts.set(qn('w:eastAsia'), CODE_FONT_NAME)
        else:
            run = p.add_run(token)

def add_rich_paragraph(text):
    """解析行内格式（**加粗**、`代码`）并添加到段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_runs_to_paragraph(p, text)
    return p

# ── 主解析循环 ─────────────────────────────────────────

i = 0
while i < len(lines):
    line = lines[i]

    # 代码块
    if line.startswith("```"):
        lang = line[3:].strip()
        i += 1
        code_lines = []
        while i < len(lines) and not lines[i].startswith("```"):
            code_lines.append(lines[i])
            i += 1
        i += 1  # 跳过结束 ```
        add_code_block(code_lines, lang)
        continue

    # 表格（检测以 | 开头的行）
    if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].startswith("|"):
            table_lines.append(lines[i])
            i += 1
        add_table_from_md(table_lines)
        continue

    # 标题
    if line.startswith("#### "):
        add_heading_styled(line[5:].strip(), 4)
        i += 1
        continue
    if line.startswith("### "):
        add_heading_styled(line[4:].strip(), 3)
        i += 1
        continue
    if line.startswith("## "):
        add_heading_styled(line[3:].strip(), 2)
        i += 1
        continue
    if line.startswith("# "):
        add_heading_styled(line[2:].strip(), 1)
        i += 1
        continue

    # 分割线
    if line.strip() == "---" or line.strip() == "---":
        doc.add_paragraph("─" * 60)
        i += 1
        continue

    # 引用块
    if line.startswith("> "):
        quote_lines = []
        while i < len(lines) and lines[i].startswith("> "):
            quote_lines.append(lines[i][2:])
            i += 1
        for ql in quote_lines:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            run = p.add_run(ql.strip())
            run.italic = True
            run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        continue

    # 无序列表
    list_match = re.match(r'^(\s*)[-*]\s+(.*)', line)
    if list_match:
        items = []
        while i < len(lines) and re.match(r'^(\s*)[-*]\s+(.*)', lines[i]):
            items.append(re.match(r'^(\s*)[-*]\s+(.*)', lines[i]).group(2))
            i += 1
        for item in items:
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            add_runs_to_paragraph(p, item)
        continue

    # 空行
    if line.strip() == "":
        i += 1
        continue

    # 普通段落
    add_rich_paragraph(line.strip())
    i += 1

# ── 保存 ────────────────────────────────────────────────
doc.save(DST)
print(f"Done! Converted to: {DST}")
print(f"File size: {os.path.getsize(DST) / 1024:.1f} KB")
